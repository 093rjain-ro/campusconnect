import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Theme colors
    BG_COLOR = RGBColor(13, 17, 23)        # #0D1117
    CARD_BG = RGBColor(22, 27, 34)         # #161B22
    BORDER_COLOR = RGBColor(48, 54, 61)    # #30363D
    CYAN = RGBColor(0, 212, 255)           # #00D4FF
    PURPLE = RGBColor(139, 92, 246)        # #8B5CF6
    TEXT_LIGHT = RGBColor(240, 246, 252)   # #F0F6FC
    TEXT_MUTED = RGBColor(139, 148, 158)   # #8B949E
    ACCENT_GREEN = RGBColor(16, 185, 129)  # #10B981

    def set_background(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = BG_COLOR

    def add_header(slide, title_text, category_text="Campus Placement System  |  Project 22"):
        set_background(slide)
        
        # Category / breadcrumb
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.4))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = CYAN

        # Slide Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.7), Inches(0.8))
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = TEXT_LIGHT

    # ─────────────────────────────────────────────────────────────
    # SLIDE 1: Title Slide
    # ─────────────────────────────────────────────────────────────
    blank_layout = prs.slide_layouts[6]
    s1 = prs.slides.add_slide(blank_layout)
    set_background(s1)

    # Accent decorative box
    accent = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.5), Inches(0.15), Inches(4.5))
    accent.fill.solid()
    accent.fill.fore_color.rgb = CYAN
    accent.line.fill.background()

    # Title box
    tb = s1.shapes.add_textbox(Inches(1.2), Inches(1.5), Inches(11), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "PROJECT 22"
    p0.font.size = Pt(18)
    p0.font.bold = True
    p0.font.color.rgb = CYAN
    p0.space_after = Pt(14)

    p1 = tf.add_paragraph()
    p1.text = "Campus Placement & \nRecruitment Drive Management System"
    p1.font.size = Pt(38)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_LIGHT
    p1.space_after = Pt(20)

    p2 = tf.add_paragraph()
    p2.text = "A relational database system for managing students, companies, recruitment drives, applications and offer workflows."
    p2.font.size = Pt(16)
    p2.font.color.rgb = TEXT_MUTED
    p2.space_after = Pt(24)

    p3 = tf.add_paragraph()
    p3.text = "DBMS Academic Project  |  End-to-End Placement Lifecycle"
    p3.font.size = Pt(13)
    p3.font.bold = True
    p3.font.color.rgb = PURPLE

    # ─────────────────────────────────────────────────────────────
    # SLIDE 2: Problem Statement
    # ─────────────────────────────────────────────────────────────
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "Problem Statement")

    intro = s2.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.7), Inches(0.5))
    p_intro = intro.text_frame.paragraphs[0]
    p_intro.text = "Most college placement cells still run on spreadsheets, emails and manual checks. That creates critical operational challenges:"
    p_intro.font.size = Pt(14)
    p_intro.font.color.rgb = TEXT_MUTED

    problems = [
        ("Ineligible students apply", "No automatic CGPA, backlog or percentage verification before student application."),
        ("Duplicate applications", "The same student can apply multiple times for the same company drive."),
        ("Rounds get lost", "Aptitude, technical and HR evaluation outcomes live in separate disconnected sheets."),
        ("Offers are hard to track", "No single source of truth for who received offers, accepted, or declined."),
        ("Reports take too long", "Compiling branch-wise or company-wise metrics requires hours of manual spreadsheet work."),
        ("Data becomes unreliable", "One accidental manual edit breaks consistency across the whole placement record set.")
    ]

    for idx, (title, desc) in enumerate(problems):
        row = idx // 2
        col = idx % 2
        left = Inches(0.8 + col * 5.9)
        top = Inches(2.2 + row * 1.55)

        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(5.6), Inches(1.4))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = BORDER_COLOR

        tf_c = card.text_frame
        tf_c.word_wrap = True
        p_t = tf_c.paragraphs[0]
        p_t.text = f"•  {title}"
        p_t.font.bold = True
        p_t.font.size = Pt(14)
        p_t.font.color.rgb = CYAN
        p_t.space_after = Pt(4)

        p_d = tf_c.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = TEXT_MUTED

    # ─────────────────────────────────────────────────────────────
    # SLIDE 3: Why This System
    # ─────────────────────────────────────────────────────────────
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "Why This System")

    why_points = [
        ("01", "Stop Ineligible Applications", "Eligibility rules (CGPA, marks, backlogs) are verified before allowing an application."),
        ("02", "One Application Per Drive", "Database-level unique constraints strictly block redundant duplicate submissions."),
        ("03", "Clear Round History", "Every Pass/Fail decision is logged in proper sequence against the correct application."),
        ("04", "Single Source for Offers", "Issue, accept, or reject final offers with one authoritative placement ledger."),
        ("05", "Instant Real-Time Reports", "Generate branch and company statistics in seconds rather than days of manual effort."),
        ("06", "Clean & Consistent Data", "Relational foreign keys and checks guarantee complete data integrity.")
    ]

    for idx, (num, title, desc) in enumerate(why_points):
        row = idx // 3
        col = idx % 3
        left = Inches(0.8 + col * 3.9)
        top = Inches(1.8 + row * 2.5)

        card = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(3.7), Inches(2.3))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = BORDER_COLOR

        tf_c = card.text_frame
        tf_c.word_wrap = True
        p_n = tf_c.paragraphs[0]
        p_n.text = num
        p_n.font.size = Pt(22)
        p_n.font.bold = True
        p_n.font.color.rgb = PURPLE
        p_n.space_after = Pt(4)

        p_t = tf_c.add_paragraph()
        p_t.text = title
        p_t.font.size = Pt(13)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_LIGHT
        p_t.space_after = Pt(6)

        p_d = tf_c.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = TEXT_MUTED

    # ─────────────────────────────────────────────────────────────
    # SLIDE 4: Objectives
    # ─────────────────────────────────────────────────────────────
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "Project Objectives")

    objectives = [
        ("1", "Identify Domain Entities & User Roles", "Map out students, departments, companies, drives, applications, rounds, and offers with tailored role permissions."),
        ("2", "Design Comprehensive ER Model", "Formulate clear cardinality, relationships, and attributes capturing the complete campus placement cycle."),
        ("3", "Relational Schema & 3NF Normalization", "Transform the conceptual ER model into normalized relational tables eliminating redundancy and anomalies."),
        ("4", "Enforce Integrity Constraints", "Implement Primary Keys, Foreign Keys, UNIQUE, and CHECK constraints at the database engine level."),
        ("5", "Sample Data & Analytical Queries", "Populate realistic student and drive data; write SQL queries for shortlist generation and placement audits."),
        ("6", "Working Application Implementation", "Deliver a functioning front-end and Python application validating end-to-end drive lifecycle operations.")
    ]

    for idx, (num, title, desc) in enumerate(objectives):
        row = idx // 2
        col = idx % 2
        left = Inches(0.8 + col * 5.9)
        top = Inches(1.8 + row * 1.7)

        card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(5.6), Inches(1.5))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = BORDER_COLOR

        tf_c = card.text_frame
        tf_c.word_wrap = True
        p_t = tf_c.paragraphs[0]
        p_t.text = f"Objective {num}: {title}"
        p_t.font.bold = True
        p_t.font.size = Pt(13)
        p_t.font.color.rgb = CYAN
        p_t.space_after = Pt(4)

        p_d = tf_c.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = TEXT_MUTED

    # ─────────────────────────────────────────────────────────────
    # SLIDE 5: Users of the System
    # ─────────────────────────────────────────────────────────────
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "Users & Responsibilities")

    users = [
        ("Placement Officer", "Create recruitment drives, define eligibility criteria, record sequential round results, issue final offers, and generate aggregate reports.", CYAN),
        ("Student / Candidate", "Explore eligible drives, submit applications, track review status across aptitude/tech/HR rounds, and accept or reject offers.", ACCENT_GREEN),
        ("Department Coordinator", "Monitor branch-wise participation, track eligible student counts, and review branch placement percentages.", PURPLE),
        ("Company Recruiter", "View drive applicant profiles, evaluate candidates across rounds, submit shortlists, and confirm hiring selections.", RGBColor(251, 191, 36))
    ]

    for idx, (role, resp, color) in enumerate(users):
        left = Inches(0.8 + idx * 2.95)
        top = Inches(1.8)

        card = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(2.8), Inches(4.8))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = BORDER_COLOR

        tf_c = card.text_frame
        tf_c.word_wrap = True
        p_r = tf_c.paragraphs[0]
        p_r.text = role
        p_r.font.size = Pt(16)
        p_r.font.bold = True
        p_r.font.color.rgb = color
        p_r.space_after = Pt(12)

        p_desc = tf_c.add_paragraph()
        p_desc.text = resp
        p_desc.font.size = Pt(12)
        p_desc.font.color.rgb = TEXT_MUTED

    # ─────────────────────────────────────────────────────────────
    # SLIDE 6: Main Functions
    # ─────────────────────────────────────────────────────────────
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "Main System Functions")

    functions = [
        ("1. Drive Setup", "Configure recruitment events tied to specific company job profiles and packages (LPA)."),
        ("2. Eligibility Check", "Automated validation of CGPA, 10th/12th percentages, and backlogs before applying."),
        ("3. Application Processing", "Single-click application with built-in prevention of duplicate submissions per drive."),
        ("4. Round Management", "Sequential progression across Aptitude, Coding, Technical, and HR interviews."),
        ("5. Offer Lifecycle", "Formal offer issuance with student acceptance/rejection logging."),
        ("6. Reporting & Analytics", "Instant shortlists, department placement audit, and recruiter analytics.")
    ]

    for idx, (title, desc) in enumerate(functions):
        row = idx // 3
        col = idx % 3
        left = Inches(0.8 + col * 3.9)
        top = Inches(1.9 + row * 2.4)

        card = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(3.7), Inches(2.1))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = BORDER_COLOR

        tf_c = card.text_frame
        tf_c.word_wrap = True
        p_t = tf_c.paragraphs[0]
        p_t.text = title
        p_t.font.bold = True
        p_t.font.size = Pt(14)
        p_t.font.color.rgb = CYAN
        p_t.space_after = Pt(6)

        p_d = tf_c.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = TEXT_MUTED

    # ─────────────────────────────────────────────────────────────
    # SLIDE 7: Process Flow
    # ─────────────────────────────────────────────────────────────
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "Process Flow & Stage Walkthrough")

    steps = [
        ("Step 1", "Create Drive", "Placement officer creates drive linked to company job profile."),
        ("Step 2", "Set Eligibility", "Defines min CGPA, max backlogs, and eligible branches."),
        ("Step 3", "Student Applies", "Eligible students submit application (blocked if ineligible or duplicate)."),
        ("Step 4", "Round Results", "Evaluation recorded sequentially: Aptitude -> Technical -> HR."),
        ("Step 5", "Issue Offer", "Offer issued to cleared candidates (one per application)."),
        ("Step 6", "Accept / Reject", "Candidate accepts or declines; placement ledger updates.")
    ]

    for idx, (step_num, title, desc) in enumerate(steps):
        left = Inches(0.8 + idx * 1.95)
        top = Inches(2.0)

        card = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(1.85), Inches(4.5))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = BORDER_COLOR

        tf_c = card.text_frame
        tf_c.word_wrap = True
        p_s = tf_c.paragraphs[0]
        p_s.text = step_num
        p_s.font.size = Pt(11)
        p_s.font.bold = True
        p_s.font.color.rgb = PURPLE
        p_s.space_after = Pt(4)

        p_t = tf_c.add_paragraph()
        p_t.text = title
        p_t.font.size = Pt(14)
        p_t.font.bold = True
        p_t.font.color.rgb = CYAN
        p_t.space_after = Pt(10)

        p_d = tf_c.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(10)
        p_d.font.color.rgb = TEXT_MUTED

    # ─────────────────────────────────────────────────────────────
    # SLIDE 8: ER Diagram (Main Entities & Cardinalities)
    # ─────────────────────────────────────────────────────────────
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "ER Diagram & Key Relationships")

    entities_str = "Department  •  Student  •  AcademicRecord  •  Skill  •  Student_Skill  •  Company  •  JobProfile  •  Drive  •  EligibilityCriteria  •  Application  •  Round  •  RoundResult  •  Offer"
    
    ent_box = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.7), Inches(0.8))
    ent_box.fill.solid()
    ent_box.fill.fore_color.rgb = CARD_BG
    ent_box.line.color.rgb = CYAN
    p_eb = ent_box.text_frame.paragraphs[0]
    p_eb.text = "13 Core Relational Entities: " + entities_str
    p_eb.font.size = Pt(11)
    p_eb.font.bold = True
    p_eb.font.color.rgb = TEXT_LIGHT

    relations = [
        ("Department  1 : N  Student", "One academic department enrolls multiple students."),
        ("Student  1 : 1  AcademicRecord", "Each student maintains exactly one official academic record."),
        ("Student  M : N  Skill", "Bridged via Student_Skill composite junction table."),
        ("Company  1 : N  JobProfile", "One hiring organization lists multiple openings."),
        ("JobProfile  1 : N  Drive", "A job profile can sponsor multiple campus drive events."),
        ("Drive  1 : 1  EligibilityCriteria", "Each recruitment drive enforces specific qualifying rules."),
        ("Student  M : N  Drive  (Application)", "Students apply to multiple drives; each drive receives multiple applications."),
        ("Drive  1 : N  Round", "Recruitment drives organize sequential evaluation rounds."),
        ("Application  1 : N  RoundResult", "Round-by-round assessment scores logged per application."),
        ("Application  1 : 1  Offer", "Successful qualification yields a unique final offer letter.")
    ]

    for idx, (rel, expl) in enumerate(relations):
        row = idx // 2
        col = idx % 2
        left = Inches(0.8 + col * 5.9)
        top = Inches(2.8 + row * 0.85)

        card = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(5.6), Inches(0.75))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = BORDER_COLOR

        tf_c = card.text_frame
        tf_c.word_wrap = True
        p_r = tf_c.paragraphs[0]
        p_r.text = rel
        p_r.font.bold = True
        p_r.font.size = Pt(11)
        p_r.font.color.rgb = CYAN

        p_e = tf_c.add_paragraph()
        p_e.text = expl
        p_e.font.size = Pt(9.5)
        p_e.font.color.rgb = TEXT_MUTED

    # ─────────────────────────────────────────────────────────────
    # SLIDE 9: Main Entities & Attributes
    # ─────────────────────────────────────────────────────────────
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "Main Entities & Attributes")

    entity_table_data = [
        ("Student", "Candidate registered for placement", "roll_no, name, email, department_id, gender"),
        ("Academic_Record", "Performance and marks overview", "student_id, cgpa, tenth_pct, twelfth_pct, backlogs"),
        ("Company", "Recruiting partner organisation", "company_id, name, industry, website, hr_email"),
        ("Job_Profile", "Specific job designation opening", "job_id, company_id, title, type, package_lpa, location"),
        ("Drive", "Campus recruitment event", "drive_id, job_id, date, venue, status"),
        ("Eligibility_Criteria", "Thresholds required to apply", "drive_id, min_cgpa, max_backlogs, min_tenth_pct, min_twelfth_pct"),
        ("Application", "Student registration for a drive", "application_id, student_id, drive_id, applied_on, status"),
        ("Round / Result", "Selection interview stage outcome", "round_id, sequence_no, application_id, result, score"),
        ("Offer", "Final formal placement job offer", "offer_id, application_id, package_lpa, status, accepted_on")
    ]

    rows = len(entity_table_data) + 1
    table_shape = s9.shapes.add_table(rows, 3, Inches(0.8), Inches(1.8), Inches(11.7), Inches(4.8))
    table = table_shape.table
    table.columns[0].width = Inches(2.2)
    table.columns[1].width = Inches(4.5)
    table.columns[2].width = Inches(5.0)

    headers = ["Entity", "Domain Concept", "Key Attributes"]
    for col_idx, h in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = CARD_BG
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = CYAN

    for row_idx, (ent, conc, attrs) in enumerate(entity_table_data):
        for col_idx, val in enumerate([ent, conc, attrs]):
            cell = table.cell(row_idx + 1, col_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = BG_COLOR
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.size = Pt(10)
            p.font.color.rgb = TEXT_LIGHT if col_idx == 0 else TEXT_MUTED
            if col_idx == 0:
                p.font.bold = True

    # ─────────────────────────────────────────────────────────────
    # SLIDE 10: Relational Schema & Constraints
    # ─────────────────────────────────────────────────────────────
    s10 = prs.slides.add_slide(blank_layout)
    add_header(s10, "Relational Schema Design")

    schema_data = [
        ("Department", "dept_id", "Referenced by Student.dept_id"),
        ("Student", "student_id", "FK -> Department(dept_id); UNIQUE(roll_no), UNIQUE(email)"),
        ("Academic_Record", "student_id", "1:1 with Student; FK -> Student(student_id)"),
        ("Skill / Student_Skill", "skill_id / composite", "M:N relationship bridge between Student and Skill"),
        ("Company", "company_id", "Referenced by Job_Profile(company_id)"),
        ("Job_Profile", "job_id", "FK -> Company(company_id)"),
        ("Drive", "drive_id", "FK -> Job_Profile(job_id)"),
        ("Eligibility_Criteria", "drive_id", "1:1 with Drive; FK -> Drive(drive_id)"),
        ("Application", "application_id", "FK -> Student, Drive; UNIQUE(student_id, drive_id)"),
        ("Round", "round_id", "FK -> Drive(drive_id); UNIQUE(drive_id, sequence_no)"),
        ("Round_Result", "result_id", "FK -> Application(app_id), Round(round_id)"),
        ("Offer", "offer_id", "FK -> Application; UNIQUE(application_id)")
    ]

    rows = len(schema_data) + 1
    table_shape10 = s10.shapes.add_table(rows, 3, Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    table10 = table_shape10.table
    table10.columns[0].width = Inches(2.4)
    table10.columns[1].width = Inches(2.2)
    table10.columns[2].width = Inches(7.1)

    headers10 = ["Table Name", "Primary Key", "Foreign Keys & Relational Constraints"]
    for col_idx, h in enumerate(headers10):
        cell = table10.cell(0, col_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = CARD_BG
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = CYAN

    for row_idx, (tbl, pk, fk) in enumerate(schema_data):
        for col_idx, val in enumerate([tbl, pk, fk]):
            cell = table10.cell(row_idx + 1, col_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = BG_COLOR
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.size = Pt(9.5)
            p.font.color.rgb = CYAN if col_idx == 0 else (PURPLE if col_idx == 1 else TEXT_MUTED)
            if col_idx < 2:
                p.font.bold = True

    # ─────────────────────────────────────────────────────────────
    # SLIDE 11: Key Business Rules Enforced
    # ─────────────────────────────────────────────────────────────
    s11 = prs.slides.add_slide(blank_layout)
    add_header(s11, "Business Rules & Enforcement Mechanism")

    rules = [
        ("One application per student per drive", "UNIQUE (student_id, drive_id) constraint on Application table."),
        ("No duplicate placement offers", "UNIQUE (application_id) constraint on Offer table."),
        ("Valid CGPA range (0.0 – 10.0)", "CHECK (cgpa >= 0.0 AND cgpa <= 10.0) constraint on Academic_Record."),
        ("Valid marks and non-negative backlogs", "CHECK (tenth_pct BETWEEN 0 AND 100) & CHECK (backlogs >= 0)."),
        ("Ordered sequential rounds", "CHECK (sequence_no > 0) + UNIQUE (drive_id, sequence_no) on Round table."),
        ("Restricted lifecycle status values", "CHECK constraints enforcing valid enum values ('Applied', 'Pass', 'Fail', 'Offered')."),
        ("Pre-application eligibility validation", "Verified in database view & application logic before allowing record insertion."),
        ("Referential integrity guarantee", "FOREIGN KEY constraints configured with appropriate ON DELETE CASCADE/RESTRICT rules.")
    ]

    for idx, (rule, enf) in enumerate(rules):
        row = idx // 2
        col = idx % 2
        left = Inches(0.8 + col * 5.9)
        top = Inches(1.8 + row * 1.25)

        card = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(5.6), Inches(1.15))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = BORDER_COLOR

        tf_c = card.text_frame
        tf_c.word_wrap = True
        p_t = tf_c.paragraphs[0]
        p_t.text = f"•  {rule}"
        p_t.font.bold = True
        p_t.font.size = Pt(12)
        p_t.font.color.rgb = CYAN
        p_t.space_after = Pt(3)

        p_e = tf_c.add_paragraph()
        p_e.text = enf
        p_e.font.size = Pt(10)
        p_e.font.color.rgb = TEXT_MUTED

    # ─────────────────────────────────────────────────────────────
    # SLIDE 12: Conclusion & Future Scope
    # ─────────────────────────────────────────────────────────────
    s12 = prs.slides.add_slide(blank_layout)
    add_header(s12, "Conclusion & Future Roadmap")

    # Left card: Current deliverables
    card_l = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.7), Inches(4.8))
    card_l.fill.solid()
    card_l.fill.fore_color.rgb = CARD_BG
    card_l.line.color.rgb = BORDER_COLOR
    tf_l = card_l.text_frame
    tf_l.word_wrap = True
    p_lt = tf_l.paragraphs[0]
    p_lt.text = "What This Project Delivers"
    p_lt.font.size = Pt(16)
    p_lt.font.bold = True
    p_lt.font.color.rgb = CYAN
    p_lt.space_after = Pt(12)

    delivs = [
        "Rigorous problem statement & user domain requirements",
        "Comprehensive 13-entity ER conceptual data model",
        "Normalized 3NF relational schema with foreign key mapping",
        "Strict business rules enforced via database constraints",
        "Live SQLite database (.db/campusconnect.db) with DB Browser sync",
        "Integrated web interface covering student, recruiter, and officer roles",
        "Complete submission-ready deliverables for Reviews 1, 2, and 3"
    ]
    for d in delivs:
        p_d = tf_l.add_paragraph()
        p_d.text = f"✔  {d}"
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = TEXT_MUTED
        p_d.space_after = Pt(4)

    # Right card: Future roadmap
    card_r = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8))
    card_r.fill.solid()
    card_r.fill.fore_color.rgb = CARD_BG
    card_r.line.color.rgb = BORDER_COLOR
    tf_r = card_r.text_frame
    tf_r.word_wrap = True
    p_rt = tf_r.paragraphs[0]
    p_rt.text = "Future Roadmap & Enhancements"
    p_rt.font.size = Pt(16)
    p_rt.font.bold = True
    p_rt.font.color.rgb = PURPLE
    p_rt.space_after = Pt(12)

    future = [
        "Automated email & SMS notifications for drive rounds and offers",
        "AI-powered resume & skill-matching engine for automated shortlists",
        "Direct integration with enterprise recruiter portals & ATS",
        "Interactive analytics dashboards with historical season trends",
        "Mobile application for push notifications and instant offer acceptance",
        "Multi-college federation support for centralized university placement"
    ]
    for f in future:
        p_f = tf_r.add_paragraph()
        p_f.text = f"🚀  {f}"
        p_f.font.size = Pt(11)
        p_f.font.color.rgb = TEXT_MUTED
        p_f.space_after = Pt(6)

    p_thanks = tf_r.add_paragraph()
    p_thanks.text = "\nThank you! Questions are welcome."
    p_thanks.font.size = Pt(13)
    p_thanks.font.bold = True
    p_thanks.font.color.rgb = TEXT_LIGHT

    out_path = os.path.join(os.path.dirname(__file__), "Project_22_Campus_Placement_System.pptx")
    prs.save(out_path)
    print(f"[SUCCESS] Presentation generated: {out_path}")

if __name__ == "__main__":
    create_deck()
